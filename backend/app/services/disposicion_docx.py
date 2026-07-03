
from pathlib import Path
import re
from tempfile import NamedTemporaryFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.disposiciones import disposicion_service
from app.services.expedientes import expediente_service
from app.services.historial import historial_service


class DisposicionDocxService:
    """Genera un DOCX institucional a partir del borrador renderizado por el motor de plantillas.

    En esta primera versión el DOCX se construye con estilos institucionales uniformes.
    La plantilla DOCX oficial queda conservada como base en storage/templates para el sprint de
    ajuste fino de estilos, encabezados y logos.
    """

    def generar_docx(self, expediente_id: str) -> Path:
        expediente = expediente_service.obtener(expediente_id)
        borrador = disposicion_service.obtener(expediente_id)

        doc = Document()
        self._configurar_documento(doc)
        self._agregar_disposicion(doc, borrador, expediente)

        out_dir = Path(__file__).resolve().parents[3] / "storage" / "exports" / expediente_id
        out_dir.mkdir(parents=True, exist_ok=True)
        safe_numero = re.sub(r"[^0-9A-Za-z_-]+", "_", expediente.numero_disposicion or expediente_id)
        salida = out_dir / f"DISPOSICION_{safe_numero}.docx"
        doc.save(salida)

        historial_service.registrar(
            expediente_id,
            "DISPOSICION_DOCX_GENERADA",
            detalle=f"Archivo generado: {salida.name}",
        )
        return salida

    def _configurar_documento(self, doc: Document) -> None:
        section = doc.sections[0]
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        styles = doc.styles
        normal = styles["Normal"]
        normal.font.name = "Times New Roman"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        normal.font.size = Pt(11)

    def _agregar_disposicion(self, doc: Document, borrador, expediente) -> None:
        texto = self._documento_completo(borrador)

        for bloque in texto.split("\n\n"):
            bloque = bloque.strip()
            if not bloque:
                continue

            if self._es_tabla_markdown(bloque):
                self._agregar_tabla_markdown(doc, bloque)
                continue

            for linea in bloque.splitlines():
                self._agregar_parrafo(doc, linea.strip())

        self._agregar_pie(doc, expediente)

    def _documento_completo(self, borrador) -> str:
        partes = [
            borrador.visto.strip(),
            borrador.considerando.strip(),
            borrador.dispone.strip(),
        ]
        return "\n\n".join(p for p in partes if p)

    def _agregar_parrafo(self, doc: Document, texto: str) -> None:
        if not texto:
            return

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.05

        upper = texto.upper().strip()
        if upper in {"CONSIDERANDO:", "DISPONE", "POR ELLO,"} or upper.startswith("EL CUERPO"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if upper.startswith("EL CUERPO") or upper == "DISPONE" else WD_ALIGN_PARAGRAPH.LEFT

        self._agregar_runs_markdown(p, texto)

    def _agregar_runs_markdown(self, paragraph, texto: str) -> None:
        # Interpreta **negrita** desde la plantilla markdown.
        partes = re.split(r"(\*\*.*?\*\*)", texto)
        for parte in partes:
            if not parte:
                continue
            bold = parte.startswith("**") and parte.endswith("**")
            contenido = parte[2:-2] if bold else parte
            run = paragraph.add_run(contenido)
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(11)
            run.bold = bold or self._debe_ir_en_negrita(contenido)

    def _debe_ir_en_negrita(self, texto: str) -> bool:
        limpio = texto.strip().upper()
        return limpio.startswith("ARTÍCULO") or limpio in {"CONSIDERANDO:", "DISPONE"}

    def _es_tabla_markdown(self, bloque: str) -> bool:
        lineas = [l for l in bloque.splitlines() if l.strip()]
        return len(lineas) >= 2 and all(l.strip().startswith("|") and l.strip().endswith("|") for l in lineas[:2])

    def _agregar_tabla_markdown(self, doc: Document, bloque: str) -> None:
        lineas = [l.strip() for l in bloque.splitlines() if l.strip()]
        filas = []
        for idx, linea in enumerate(lineas):
            if idx == 1 and set(linea.replace("|", "").replace(":", "").replace(" ", "")) <= {"-"}:
                continue
            celdas = [c.strip() for c in linea.strip("|").split("|")]
            filas.append(celdas)

        if not filas:
            return

        tabla = doc.add_table(rows=1, cols=len(filas[0]))
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabla.style = "Table Grid"
        self._set_table_borders(tabla, visible=True)

        hdr = tabla.rows[0].cells
        for i, val in enumerate(filas[0]):
            self._set_cell_text(hdr[i], val, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        for fila in filas[1:]:
            cells = tabla.add_row().cells
            for i, val in enumerate(fila):
                align = WD_ALIGN_PARAGRAPH.RIGHT if i == len(fila) - 1 else WD_ALIGN_PARAGRAPH.LEFT
                bold = "**" in val or "Monto Neto a Pagar" in val
                val = val.replace("**", "")
                self._set_cell_text(cells[i], val, bold=bold, align=align)

        doc.add_paragraph()

    def _set_cell_text(self, cell, text: str, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(10.5)
        run.bold = bold

    def _set_table_borders(self, table, visible: bool = True) -> None:
        tbl = table._tbl
        tblPr = tbl.tblPr
        borders = tblPr.first_child_found_in("w:tblBorders")
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tblPr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            tag = "w:" + edge
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            element.set(qn("w:val"), "single" if visible else "nil")
            element.set(qn("w:sz"), "4")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "auto")

    def _agregar_pie(self, doc: Document, expediente) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(f"DISPOSICION N° {expediente.numero_disposicion or '____/____'}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.bold = True

        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run2 = p2.add_run(f"EXPEDIENTE: {expediente.numero_interno}")
        run2.font.name = "Times New Roman"
        run2.font.size = Pt(11)
        run2.bold = True


disposicion_docx_service = DisposicionDocxService()
